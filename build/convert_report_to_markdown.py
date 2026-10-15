from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'output' / 'report' / 'AI_Economic_Perceptions_Consumer_Sentiment_Report_v4.docx'
TARGET = ROOT / 'output' / 'report' / 'AI_Economic_Perceptions_Consumer_Sentiment_Report.md'

def clean(value: str) -> str:
    return ' '.join(value.replace('|', '\\|').split())

def heading_level(style: str) -> int | None:
    if style == 'Title': return 1
    if style.startswith('Heading '):
        try: return min(int(style.split()[-1]) + 1, 6)
        except ValueError: return None
    return None

def emit_table(table) -> list[str]:
    rows = [[clean(cell.text) for cell in row.cells] for row in table.rows]
    if not rows: return []
    width = len(rows[0])
    rows = [r + [''] * (width - len(r)) for r in rows]
    lines = ['| ' + ' | '.join(rows[0]) + ' |', '| ' + ' | '.join(['---'] * width) + ' |']
    lines += ['| ' + ' | '.join(row) + ' |' for row in rows[1:]]
    return lines

def main():
    doc = Document(SOURCE)
    table_by_element = {table._element: table for table in doc.tables}
    lines = []
    for element in doc.element.body.iterchildren():
        if element.tag.endswith('}p'):
            paragraph = next((p for p in doc.paragraphs if p._p is element), None)
            if paragraph is None: continue
            value = clean(paragraph.text)
            if not value: continue
            style = paragraph.style.name
            level = heading_level(style)
            if level:
                lines.extend(['#' * level + ' ' + value, ''])
            elif style.startswith('List Bullet'):
                lines.append(f'- {value}')
            elif style.startswith('List Number'):
                lines.append(f'1. {value}')
            else:
                lines.extend([value, ''])
        elif element.tag.endswith('}tbl'):
            table = table_by_element.get(element)
            if table:
                lines.extend(emit_table(table)); lines.append('')
    # Normalize excess blank lines while retaining readable Markdown structure.
    normalized=[]
    for line in lines:
        if line == '' and normalized and normalized[-1] == '': continue
        normalized.append(line)
    TARGET.write_text('\n'.join(normalized).strip() + '\n', encoding='utf-8')
    print(TARGET)

if __name__ == '__main__': main()
