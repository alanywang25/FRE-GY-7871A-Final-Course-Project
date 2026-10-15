from pathlib import Path
from datetime import date
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output' / 'report' / 'AI_Economic_Perceptions_Consumer_Sentiment_Report_v4.docx'
ASSETS = ROOT / 'build' / 'report_assets'
ASSETS.mkdir(parents=True, exist_ok=True)

NAVY = '17365D'; BLUE = '2F75B5'; LIGHT = 'D9EAF7'; GRID = 'D9D9D9'; GRAY = '595959'

def set_cell(cell, fill=None, color='000000', bold=False, size=8.5):
    if fill:
        tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        for r in p.runs:
            r.font.color.rgb = RGBColor.from_string(color); r.font.bold = bold; r.font.size = Pt(size)

def borders(table):
    tblPr = table._tbl.tblPr
    el = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement(f'w:{edge}'); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6'); e.set(qn('w:color'),GRID); el.append(e)
    tblPr.append(el)

def shade_rows(table):
    for j, cell in enumerate(table.rows[0].cells): set_cell(cell, NAVY, 'FFFFFF', True, 8.5)
    for i,row in enumerate(table.rows[1:],1):
        for cell in row.cells: set_cell(cell, 'F5F9FC' if i%2==0 else 'FFFFFF')

def add_table(doc, headers, rows, widths=None):
    table=doc.add_table(rows=1, cols=len(headers)); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'
    for i,h in enumerate(headers): table.rows[0].cells[i].text=str(h)
    for row in rows:
        cells=table.add_row().cells
        for i,val in enumerate(row): cells[i].text=str(val)
    if widths:
        for row in table.rows:
            for c,w in zip(row.cells,widths): c.width=Inches(w)
    borders(table); shade_rows(table)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return table

def add_heading(doc, text, level=1):
    p=doc.add_paragraph(style=f'Heading {level}'); p.add_run(text)
    p.paragraph_format.space_before=Pt(8 if level==1 else 5); p.paragraph_format.space_after=Pt(5)
    return p

def add_para(doc, text, bold_lead=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.08
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold=True; p.add_run(text[len(bold_lead):])
    else: p.add_run(text)
    return p

def add_bullets(doc, items):
    for x in items:
        p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.02; p.add_run(x)

def caption(doc, text):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(6)
    r=p.add_run(text); r.italic=True; r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY)

def page(doc): doc.add_page_break()

def style(doc):
    sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.62); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
    styles=doc.styles
    styles['Normal'].font.name='Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Aptos'); styles['Normal'].font.size=Pt(9.5)
    for name,size in [('Title',24),('Subtitle',12),('Heading 1',15),('Heading 2',11)]:
        st=styles[name]; st.font.name='Aptos'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Aptos'); st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0)
    styles['Heading 1'].font.bold=True; styles['Heading 2'].font.bold=True
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=footer.add_run('AI Applications in Finance  |  Draft report  |  October 2026'); rr.font.size=Pt(8); rr.font.color.rgb=RGBColor.from_string(GRAY)

def main():
    doc=Document(); style(doc)
    # Cover
    p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(130); p.add_run('AI Economic Perceptions and Consumer Sentiment')
    p=doc.add_paragraph(style='Subtitle'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('A measurement and forecasting framework ahead of the 2026 midterm elections')
    doc.add_paragraph('')
    for t in ['Course: AI Applications in Finance','Final project report draft','Analysis cutoff: October 14, 2026']:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(t).font.size=Pt(11)
    page(doc)
    # Executive summary
    add_heading(doc,'Executive summary')
    add_para(doc,'This project develops a reproducible framework for separating economic conditions, economic narratives, and consumer sentiment. It asks whether an economic-only model leaves a persistent Michigan consumer-sentiment gap and whether carefully measured media variables improve understanding of that gap. The project treats a negative residual as unexplained pessimism under a specified model, not as proof that consumers are irrational or that coverage is biased.')
    add_heading(doc,'Findings',2)
    add_bullets(doc,[
        'A chronology-respecting economic-only rolling forecast has an out-of-sample MAE of 12.579 index points and RMSE of 16.458 across 126 forecasts. Its residuals are strongly autocorrelated, so it is a baseline rather than a complete explanation.',
        'The project adds a household-stress proxy and a broader control set. The expanded-controls forecast reduces MAE to 4.467 and RMSE to 6.236 across 142 out-of-sample forecasts, demonstrating that headline macro indicators are not the whole household experience.',
        'Two source-blinded human coders independently labeled 312 retained items. Tone agreement was 59.3% with Cohen kappa of 0.437; disagreements were adjudicated before a monthly media measure was constructed.',
        'In the 2014-2021 historical archive, a 10 percentage-point increase in prior-month negative tone is associated with an 8.09-point lower headline sentiment reading after controls. This is an observational association, not a causal media effect.',
        'Adding historical media tone does not improve MAE in the primary historical rolling forecast (6.04 versus 5.93 for economic-only), even though RMSE declines modestly (9.84 versus 10.63). The practical conclusion is to monitor narratives as risk signals, not to insert them mechanically into a production forecast.'
    ])
    add_heading(doc,'Practical implication',2)
    add_para(doc,'For financial forecasting and risk management, sentiment should be monitored as a potentially informative but noisy state variable. A model-risk process should retain real-time vintages, track the sentiment gap and media-measure quality, and escalate large unexplained gaps for human review rather than labeling them “bias.”')
    # question/design
    add_heading(doc,'Research question and analytical design')
    add_para(doc,'The project evaluates three competing explanations for sentiment that appears weak relative to selected macro indicators: unmeasured household pressure, differences in information and framing, and measurement or timing limitations. The design therefore separates prediction from explanation and association from causal impact.')
    add_table(doc,['Question','Operational response','What the evidence can support'],[
        ['Economic baseline','Monthly, availability-aware rolling forecasts using U-3, PCE inflation, and quarterly real GDP growth.','Out-of-sample predictive performance and a sentiment gap.'],
        ['Household pressures','Supplementary controls include U-6, real disposable income, cumulative PCE-price change, debt service, housing affordability, and policy uncertainty.','Whether broader aggregates improve forecasts; not household-level causal effects.'],
        ['Narrative measurement','Source-blinded coding of tone, salience, framing, factual calibration, and political attribution.','Descriptive content-supply measures and selected cross-source contrasts.'],
        ['Media and sentiment','Chronological media models and HAC association estimates with a future-media placebo.','Associations and forecasting comparisons, not media-exposure effects.'],
        ['Causal evidence','Pre-specified randomized truthful-framing experiment proposed for future work.','A feasible next step, subject to approval and fielding.']
    ],[1.3,2.25,3.1])
    add_heading(doc,'Pre-registered interpretation rules',2)
    add_bullets(doc,['Negative coverage is not defined as biased merely because it is negative. Bias would require a declared benchmark, such as comparable releases or matched topic-day content.', 'Text data can measure content supply and framing but cannot establish what Michigan respondents encountered, recalled, or believed.', 'A null or weak media forecast improvement would be evidence against a strong incremental-prediction claim; it would not prove that media never matters.'])
    # baseline/data
    add_heading(doc,'Economic and survey baseline')
    add_para(doc,'The analysis begins in 2010 at monthly frequency. The supplied course workbook provides the Michigan headline index, U-3 unemployment, PCE inflation, and quarterly real GDP growth. The reproducible package supplements it with documented public macro series and an ALFRED-vintage archive. Predictions use conservative availability rules and do not interpolate quarterly GDP into artificial monthly observations.')
    add_table(doc,['Series','Role and timing treatment','Key interpretation safeguard'],[
        ['Michigan consumer sentiment','Outcome; headline index plus available Current Conditions and Expectations subindices.','The survey covers financial, business, and buying views; it is not a direct vote on U-3, PCE, or GDP.'],
        ['U-3 unemployment','Core economic predictor; seasonally adjusted official rate.','Labor-force definition differs from broader household stress and underemployment.'],
        ['PCE inflation','Year-over-year price-index change.','Disinflation does not reverse the level effect of previously elevated prices.'],
        ['Real GDP growth','Quarterly quarter-over-quarter annualized growth; held at its actual release cadence.','Aggregate output growth can diverge from household experience.'],
        ['Supplementary controls','U-6, real income, cumulative prices, debt service, housing affordability, policy uncertainty.','Controls improve description; they do not solve causal identification.']
    ],[1.2,2.5,2.95])
    add_table(doc,['Latest stored month','Observed sentiment','Economic-only prediction','Sentiment gap','95% prediction interval'],[['October 2026 (partial period)','46.3','75.3','-29.0','[51.9, 98.7]']],[1.25,1.2,1.45,1.0,1.75])
    caption(doc,'Table 2. Latest stored baseline comparison. October is explicitly treated as a partial-period observation. Source: project output economic_only_oos.csv.')
    add_para(doc,'The latest stored observation is partial-period October 2026 and should not be treated as a completed month. The dashboard visibly distinguishes data quality and availability from inferred conditions.')
    # corpus
    add_heading(doc,'Balanced media corpus and measurement framework')
    add_para(doc,'The corpus is a documented content-supply sample, not a media-exposure panel. It combines economic news cards/headlines from 16 national and local publishers with a public Mastodon supplement. The collection records query, timestamp, source, retained text, deduplication, and automation flags. It does not scrape paywalled article bodies, comments, advertisements, or profiles.')
    add_table(doc,['Dimension','Operationalization in the project','Boundary on interpretation'],[
        ['Tone','Negative, neutral, positive, mixed, or unclear economic language; separate factual description from emotional wording.','Negative tone is not equivalent to bias.'],
        ['Salience and selection','Mentions of layoffs, hiring, prices, disinflation, growth, and recession risk.','Supply does not equal audience attention.'],
        ['Framing','Price levels versus inflation rates; household experience versus national aggregates; current conditions versus forecasts.','Framing can be described without inferring a reader cognitive bias.'],
        ['Factual calibration','Numerical claims assessed against information available at publication.','Incomplete or unverifiable items are retained as such.'],
        ['Political attribution','Credit or blame assigned to institutions or actors.','Criticism is not automatically bias.']
    ],[1.25,3.1,2.3])
    add_para(doc,'A matched source-comparative benchmark pairs items within publication-day and pre-specified topic cells. Across 18 local-national matched cells, national items had a 21.3 percentage-point higher negative-tone share, but the 95% normal-approximation interval spans -1.24 to 43.83 points. The evidence is too imprecise to support a broad comparative claim.')
    # validation
    add_heading(doc,'AI-assisted measurement validation')
    add_para(doc,'The project uses machine labels only as triage and benchmarked classifier outputs, not as ground truth or a second “human” rater. Two independent human annotators received randomized, source-blinded packets; model labels are excluded from human-agreement statistics by design. The completed all-field workflow covers 312 finalized items.')
    add_table(doc,['Validation metric','Result','Interpretation'],[
        ['Paired independent human tone labels','312','Exceeds the course minimum of 300.'],
        ['Raw tone agreement','59.3%','Annotators often differ on nuanced economic language.'],
        ['Cohen kappa','0.437','Moderate agreement; supports use with adjudication and disclosed uncertainty.'],
        ['Adjudication','Completed before aggregation','Original votes and decision rationale are preserved.'],
        ['Invariance checks','Actor-name/source masking protocol included','A no-change dictionary result is not a validation claim for an LLM.']
    ],[2.25,1.25,3.15])
    add_table(doc,['Coded month','Items','Negative tone','Positive tone','Mixed tone','Sources'],[['Aug 2026','28','32.1%','28.6%','14.3%','10'],['Sep 2026','69','29.0%','17.4%','15.9%','10'],['Oct 2026','215','35.8%','14.9%','7.4%','17']],[1.1,.65,1.0,1.0,.8,.75])
    caption(doc,'Table 4. Adjudicated 2026 media-tone mix. Availability reflects final-label build timing, not publication-date exposure. Source: project media_monthly.csv.')
    add_para(doc,'The measurement process preserves item IDs, exact retained text, source metadata, evidence spans, prompt/model records, adjudication records, and uncertainty. This makes future audit and recalibration possible while keeping the substantive claims limited to the measured construct.')
    # forecast
    add_heading(doc,'Forecasting results and model comparison')
    add_para(doc,'The core forecast comparison uses chronological holdouts. Reported metrics are out-of-sample and reflect the exact information-contract implemented in the pipeline. Performance should be read alongside residual diagnostics, because a low error can still conceal serial dependence or timing leakage.')
    add_table(doc,['Model','Out-of-sample n','MAE','RMSE','95% interval coverage','Residual lag-1 autocorrelation'],[
        ['Economic-only macro baseline','126','12.579','16.458','73.8%','0.830'],
        ['Economic + lagged sentiment sensitivity','141','3.451','4.585','90.1%','0.166'],
        ['Household-stress proxy extension','142','8.178','12.062','71.8%','0.623'],
        ['Expanded macro controls','142','4.467','6.236','83.1%','0.586']
    ],[2.25,1.0,.55,.55,1.15,1.0])
    add_para(doc,'The lagged-sentiment sensitivity model is useful diagnostically, but it should not be represented as a causal model or an operationally clean explanation of sentiment; it directly leverages outcome persistence. The expanded-control result is more relevant to the central research question: omitted household-relevant variables can materially improve a macro-only baseline.')
    add_heading(doc,'Historical media-model comparison',2)
    add_table(doc,['Historical specification, 2014-2021','n','MAE','RMSE','Interpretation'],[
        ['Economic-only','41','5.929','10.629','Reference model on common historical archive.'],
        ['Economic + one-month media tone','41','6.040','9.836','No MAE improvement; modest RMSE reduction is not sufficient for a strong predictive claim.'],
        ['Economic + media + archived controls','41','7.920','18.612','Worse out-of-sample performance; richer controls can increase estimation risk in a short sample.']
    ],[2.9,.45,.6,.65,2.05])
    add_heading(doc,'How the competing hypotheses fare',2)
    add_para(doc,'The omitted-household-pressure hypothesis receives the strongest forecasting support: a broader control set materially improves error relative to the narrow baseline. The information-and-framing hypothesis remains plausible but unproven: tone is measurable and associated with sentiment, yet it does not improve common-sample MAE and the design cannot distinguish exposure from common shocks or demand. The timing-and-model-limit hypothesis is also supported by the baseline residual autocorrelation, release-vintage sensitivity, survey-method caveats, and the partial-period cutoff.')
    # association
    add_heading(doc,'Observational association and causal interpretation')
    add_para(doc,'The historical association model is deliberately labeled associational. It uses the prior-calendar-month negative-tone share, lagged outcome, unemployment, PCE inflation, quarterly GDP growth, calendar seasonality, and HAC(6) uncertainty. It also includes a future-media placebo as a reverse-timing diagnostic. The archive contains 89 aligned months from January 2014 through June 2021.')
    add_table(doc,['Outcome','Coefficient per 10pp','HAC 95% CI','Future-media placebo p-value'],[
        ['Headline sentiment','-8.09','[-11.73, -4.45]','0.543'],
        ['Current conditions','-8.44','[-16.78, -0.10]','0.995'],
        ['Expectations','-7.31','[-8.83, -5.79]','0.529']
    ],[1.8,1.35,1.65,1.7])
    add_para(doc,'These coefficients do not identify an effect of media tone. Economic shocks, political events, publisher and audience demand, selective consumption, and unobserved household conditions can jointly affect coverage and sentiment. A nonzero future-media placebo would count against a simple directional account; a null placebo does not demonstrate causality.')
    add_heading(doc,'Feasible causal next step',2)
    add_para(doc,'The documented randomized truthful-framing experiment would assign consenting respondents to balanced factual level-and-change framing, an uncertainty framing using the same verified facts, or a neutral factual control. It would pre-specify a primary expectations outcome, report intention-to-treat estimates with robust uncertainty, and require instructor or IRB approval before fielding. It would estimate effects in the recruited sample, not automatically on the national Michigan index.')
    # dashboard/risk
    add_heading(doc,'Monitoring dashboard and financial forecasting implication')
    add_para(doc,'The static dashboard is an updateable monitoring surface rather than a production trading signal. It displays sentiment, economic-model prediction and uncertainty, the sentiment gap, source composition, media measurement status, macro refresh status, model metrics, and data-quality warnings. Empty media values remain missing rather than being converted to zeros.')
    add_table(doc,['Risk-management use','Recommended operating rule','Why it is appropriate'],[
        ['Macro scenario review','Flag sustained, large negative sentiment gaps for analyst review.','A gap may reflect omitted household pressure, timing, or narratives; it is a diagnostic, not a causal label.'],
        ['Consumer-sensitive forecasting','Maintain a narrative-risk overlay alongside real-time economic vintages.','Narrative measures may help characterize uncertainty even when incremental MAE does not improve.'],
        ['Model governance','Version data, prompts, coding decisions, release timing, and model outputs.','The historical results are sensitive to short samples and availability assumptions.'],
        ['Communication risk','Describe media measures as content supply and report missingness.','Avoids implying exposure, persuasion, or ideology that the data cannot observe.']
    ],[1.55,2.65,2.55])
    add_para(doc,'The practical implication is conservative: risk teams should monitor discrepancies between aggregate macro conditions and household perceptions, but use corroborating data before changing forecasts, asset allocation, credit policy, or public narratives. The current evidence supports a structured review process—not a claim that negative media coverage caused the 2026 sentiment gap.')
    # limitations
    add_heading(doc,'Limitations, ethics, and prospective update plan')
    add_heading(doc,'Material limitations',2)
    add_bullets(doc,[
        'The original workbook has only the headline Michigan index, U-3, PCE YoY, and quarterly GDP growth. Subindices and added controls have different availability and revision characteristics.',
        'The historical media archive has 89 aligned months. Coefficients are unstable across specifications, and no respondent-level exposure or political-affiliation measure is available.',
        'The 2026 media aggregate is finalized after human coding. It is descriptive and should not be shown as a contemporaneously available predictor without rebuilding it under a prospective timestamp contract.',
        'Michigan subindices are final-report backfills for historical association, not evidence of real-time historical availability. Survey methodology changes also limit naïve comparisons across periods.',
        'The corpus samples selected headlines/news cards and public social posts, not all articles or all household information environments.'
    ])
    add_heading(doc,'Ethics and deployment safeguards',2)
    add_para(doc,'The project uses public or appropriately licensed material, removes social-account identifiers from retained outputs, and requires consent and approval before recruiting human participants. It avoids treating political criticism as bias, refrains from inferring individual cognition from text, and preserves data provenance and model artifacts for audit.')
    add_heading(doc,'Prospective update plan',2)
    add_para(doc,'Before the election, future updates should freeze a dated vintage of every macro input, record release times, preserve media collection and coding availability timestamps, rerun all validation gates, publish missingness, and version the dashboard payload. A monitoring update should report the observation cutoff and should not attribute movement to election proximity merely because it occurs during the campaign.')
    # conclusion
    add_heading(doc,'Conclusion')
    add_para(doc,'The project delivers a reproducible framework for analyzing economic perceptions without assuming that pessimism is irrational or that negative coverage is biased. The strongest evidence is that a narrow macro baseline leaves meaningful residual variation and that broader household-relevant controls can improve predictive accuracy. Human-coded media measures can be built and audited, but the present historical results do not justify a causal media claim.')
    add_para(doc,'For finance, the durable contribution is a governance approach: treat consumer sentiment and measured narratives as monitored risk indicators; keep timing, revisions, and measurement uncertainty explicit; and escalate unexplained gaps for research rather than converting them into strong causal stories.')
    add_heading(doc,'References and reproducible materials')
    refs=[
        '[1] University of Michigan, Surveys of Consumers. Consumer Sentiment Index and methodological documentation.',
        '[2] U.S. Bureau of Labor Statistics. Labor Force Statistics from the Current Population Survey, unemployment rate (U-3).',
        '[3] U.S. Bureau of Economic Analysis. Personal Consumption Expenditures Price Index and Gross Domestic Product releases.',
        '[4] Federal Reserve Bank of St. Louis, ALFRED. Vintage macroeconomic data used for point-in-time controls.',
        '[5] Baker, Bloom, and Davis. Economic Policy Uncertainty index documentation.',
        '[6] Project repository: README.md; data/README.md; data/corpus/README.md; docs/ai_measurement_validation.md; docs/causal_design_protocol.md; data/processed/run_summary.json; data/processed/media_model_status.json; data/processed/causal_association_results.csv.',
        '[7] Course brief: Final Project Applications in Finance, October 2026.'
    ]
    for x in refs: add_para(doc,x)
    add_para(doc,'Drafting note: This report is a working draft based on the current repository as of October 14, 2026. Confirm course-specific authorship and AI-use disclosure requirements before submission.')
    # Single major transition before the embedded evidence inventory; sections otherwise flow continuously.
    page(doc)
    # Main-report evidence inventory: integrates substantive notebook and dashboard analyses.
    add_heading(doc,'Dashboard inventory and current monitoring state')
    add_para(doc,'This section maps every substantive dashboard display to its corresponding report evidence. The dashboard is a monitoring surface: it presents values, availability, and diagnostics without converting missing measures to zeros or interpreting a residual as a causal media effect.')
    add_table(doc,['Dashboard display','Included analysis or current state','Interpretation boundary'],[
        ['Latest indicator strip','October 2026 partial-period sentiment 46.3; economic-only prediction 75.3; gap -29.0; 95% interval [51.9, 98.7].','Partial-period status is explicit; a gap is not proof of bias.'],
        ['Forecast and gap charts','126 availability-aware out-of-sample forecast observations; observed sentiment, prediction, interval, and residual.','Uses release-date ALFRED vintages; date-level chronology, not intraday ordering.'],
        ['Macro context chart','Available monthly unemployment series, alongside documented PCE and GDP inputs.','GDP remains quarterly; the pipeline does not manufacture monthly GDP observations.'],
        ['Holdout scorecard','Economic-only: n=126, MAE 12.579, RMSE 16.458, coverage 73.8%.','Historical holdout performance, not a causal estimate.'],
        ['Corpus readiness','539 collected records; 315 balanced coding sample; 105 each national, local, and social; 312 paired labels.','Source-blinded aggregate is descriptive, not audience exposure.'],
        ['Validated media panels','Monthly tone, factual calibration, salience, framing, source composition, item count, and missingness.','Final-label availability time is not a real-time publication timestamp.']
    ],[1.55,3.15,1.8])
    add_table(doc,['Additional dashboard display','Included analysis or current state'],[
        ['Source-comparative benchmark','171 matched items across 25 day-topic cells and 21 source-measure contrasts; content-supply comparison only.'],
        ['Age-group ICS chart','Published three-month moving average by age; displayed cases are source metadata, not pooled analytic bases.'],
        ['Michigan components','Current Conditions and Expectations final-value history with date-level final-release timing.'],
        ['Model ladder and evidence readiness','Seven estimated forecast variants; validation, adjudication, classifier, historical-media, vintage, and refresh readiness states.'],
        ['Causal-analysis panel','89-observation retrospective association, HAC uncertainty, future-media placebo, and explicit non-identification statement.'],
        ['Release chronology and caveats','Verified final survey releases, monitoring cutoff, method-comparability warning, and availability limitations.']
    ],[2.05,4.45])
    add_heading(doc,'Reproducible notebook outputs')
    add_para(doc,'The notebook is a reproducible narrative around the pipeline. Its code cells construct the following tables and diagnostics; the notebook itself does not retain rendered plot outputs, so the report reproduces the results from the saved pipeline artifacts and dashboard payload.')
    add_table(doc,['Notebook section','Output or analysis','Included evidence'],[
        ['Data provenance and timing','Non-missing counts and first/last observation dates for sentiment, U-3, PCE YoY, and GDP growth.','Data-contract discussion in Section 3; appendix inventory and refresh status.'],
        ['Economic-only rolling OLS','Metrics, final 12 predictions with intervals, interval coverage, and 10 largest negative gaps.','Sections 3 and 7; latest stored forecast table; dashboard forecast and gap displays.'],
        ['Autocorrelation and vintages','Lagged-sentiment sensitivity and month-end ALFRED initial-vintage comparator.','Model-comparison table and residual-autocorrelation discussion.'],
        ['Household-stress extension','Incremental forecast comparison using unemployment change and GDP recession signal.','Forecasting results table and interpretation.'],
        ['Expanded public controls','Refresh manifest plus available control columns and current-vintage guardrail.','Data provenance, model comparison, and evidence-readiness inventory.'],
        ['Corpus and validation','Corpus status, independent labels, adjudication, dictionary/classifier diagnostic, and invariance test.','Measurement and dashboard sections of this report.'],
        ['Causal association','Status, coverage, non-identification statement, and results table for three sentiment outcomes.','Association and causal-design sections of this report.'],
        ['Historical media and subgroups','Monthly media aggregate, media-model ladder, subindex and subgroup contracts.','Forecasting, monitoring, and survey sections of this report.'],
        ['Dashboard build','Normalized panel, out-of-sample outputs, metadata, and current payload.','Dashboard evidence tables in this report.']
    ],[1.65,2.5,2.35])
    add_heading(doc,'Expanded model ladder',2)
    add_table(doc,['Model','n','MAE','RMSE','Coverage'],[
        ['Economic-only, release-date ALFRED vintages','126','12.579','16.458','73.8%'],
        ['Economic-only, assumed-lag sensitivity','142','7.665','11.629','74.6%'],
        ['Economic + lagged sentiment sensitivity','141','3.451','4.585','90.1%'],
        ['Economic + household-stress proxy','142','8.178','12.062','71.8%'],
        ['Economic + expanded public controls','142','4.467','6.236','83.1%'],
        ['Economic + historical media','41','6.040','9.836','70.7%'],
        ['Economic + media + archived controls','41','7.920','18.612','68.3%']
    ],[3.0,.45,.65,.65,.95])
    add_heading(doc,'Complete measurement and benchmark summary')
    add_heading(doc,'Current media measures and source composition',2)
    add_table(doc,['Month','Items','Sources','Negative','Positive','Mixed','Calibrated','Price-level frame','Household frame'],[
        ['Aug 2026','28','10','32.1%','28.6%','14.3%','100.0%','21.4%','28.6%'],
        ['Sep 2026','69','10','29.0%','17.4%','15.9%','98.6%','14.5%','21.7%'],
        ['Oct 2026','215','17','35.8%','14.9%','7.4%','85.1%','14.9%','18.6%']
    ],[.7,.52,.58,.7,.7,.6,.8,1.0,1.0])
    add_heading(doc,'Matched content-supply benchmark',2)
    add_table(doc,['Measure and comparison','Matched cells','Difference B minus A','95% normal-approximation interval','Interpretation'],[
        ['Negative tone: national minus local','18','+21.30 pp','[-1.24, 43.83]','Imprecise; not an exposure or causal estimate.'],
        ['Positive tone: national minus local','18','-13.43 pp','[-46.23, 19.38]','Imprecise; not an exposure or causal estimate.'],
        ['Materially incomplete calibration: national minus local','18','0.00 pp','[0.00, 0.00]','No contrast in this matched subset.'],
        ['Materially incomplete calibration: social minus national','10','+19.63 pp','[0.60, 38.67]','Content-supply contrast only; small matched-cell count.'],
        ['Political attribution: national minus local','18','+11.11 pp','[-9.17, 31.40]','Imprecise; not ideological-bias evidence.']
    ],[2.2,.75,.9,1.25,1.55])
    add_heading(doc,'Classifier and invariance diagnostics',2)
    add_table(doc,['Diagnostic','Result','Boundary'],[
        ['Dictionary tone baseline','Accuracy 0.580; macro-F1 0.483 on 312 adjudicated labels.','Development-set diagnostic; not a final production-model performance claim.'],
        ['Per-class tone results','Mixed F1 0.215; negative 0.415; neutral 0.695; positive 0.605.','Performance varies materially by class.'],
        ['Actor/outlet masking','0 actor and 0 source label changes across 315 dictionary items.','Diagnostic only; does not demonstrate fairness or LLM invariance.'],
        ['Full-field classifier','252 development and 60 released holdout rows; no further tuning permitted.','Held-out diagnostic only; preserve model hash and run baseline.']
    ],[1.6,2.65,2.4])
    add_heading(doc,'Survey, causal, and release-monitoring analyses')
    add_heading(doc,'Causal-association results and falsification check',2)
    add_table(doc,['Outcome','n','Coefficient per 10pp','HAC SE','95% CI','R squared','Future-tone placebo p'],[
        ['Headline sentiment','89','-8.091','1.859','[-11.735, -4.448]','0.825','0.543'],
        ['Current Conditions','89','-8.440','4.256','[-16.781, -0.098]','0.826','0.995'],
        ['Expectations','89','-7.312','0.775','[-8.831, -5.794]','0.730','0.529']
    ],[1.2,.45,1.0,.7,1.15,.7,1.0])
    add_para(doc,'The model is retrospective and associational. It cannot identify respondent exposure, selection into consumption, changes in individual beliefs, a causal effect on the national Michigan index, or confounding from economic shocks, political events, publisher/audience demand, and unmeasured household conditions.')
    add_heading(doc,'Survey components, subgroups, and release chronology',2)
    add_table(doc,['Dashboard analysis','Coverage or display','Interpretation limit'],[
        ['Current Conditions and Expectations','200 final-value records; historical annual-report backfills for 2010–2025.','Usable for retrospective association; not a real-time historical-availability claim.'],
        ['Published age-group ICS','600 dashboard records; published three-month moving averages.','Cases are source metadata; not pooled analytic bases.'],
        ['Verified final survey releases','36 dated Michigan final releases from Sep 2023 through Aug 2026.','Dates are verified; no intraday release time is asserted.'],
        ['Monitoring update plan','Refresh after each verified Michigan final release and major-event entry through election day.','Archive inputs, vintages, labels, and outputs for every run.']
    ],[1.75,2.4,2.5])
    add_heading(doc,'Interpretive caveats retained in the dashboard',2)
    add_bullets(doc,['Michigan’s 2024 transition from telephone random-digit dialing to address-based web collection requires a documented comparability sensitivity before interpreting cross-transition changes.', 'The public macro refresh is current-vintage monitoring data; historical real-time testing requires archived release vintages.', 'The dashboard does not assert verified subgroup ICC/ICE coverage after February 2026 because the March–August archive query returned no records.', 'Any prospective analysis must keep the partial-month policy, date cutoff, release documentation, and missing-data flags visible.'])
    OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT); print(OUT)

if __name__=='__main__': main()
